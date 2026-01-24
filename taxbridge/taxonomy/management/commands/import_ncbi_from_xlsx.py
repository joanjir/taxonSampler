# taxbridge/taxonomy/management/commands/import_ncbi_from_xlsx.py
from __future__ import annotations

from pathlib import Path
import openpyxl

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from taxonomy.models import Taxon
from taxonomy.services.ncbi_taxonomy import resolve_name_to_taxon


class Command(BaseCommand):
    help = "Importa taxones NCBI desde XLSX (columna Organism). Guarda mínimo: taxid + scientific_name."

    def add_arguments(self, parser):
        parser.add_argument(
            "--file",
            required=True,
            help="Ruta al XLSX (ej: C:\\...\\taxbridge\\data\\metazoa_genomes_taxonomy.xlsx)",
        )
        parser.add_argument("--sheet", default="Sheet1", help="Nombre de hoja (default: Sheet1)")
        parser.add_argument("--name-col", default="Organism", help="Columna con el organismo (default: Organism)")

        parser.add_argument("--limit", type=int, default=0, help="0 = sin límite (filas a leer)")
        parser.add_argument("--max-new", type=int, default=0, help="0 = sin límite (nuevos reales a insertar)")

        parser.add_argument(
            "--update-existing",
            action="store_true",
            help="Si se activa, actualiza scientific_name/rank en existentes; si no, los omite.",
        )

        parser.add_argument("--dry-run", action="store_true", help="No escribe en DB, solo reporta.")
        parser.add_argument("--log-file", default="", help="Archivo .log opcional.")

    def handle(self, *args, **opts):
        xlsx = Path(opts["file"])
        if not xlsx.exists():
            raise CommandError(f"No existe el archivo: {xlsx}")

        sheet_name = opts["sheet"]
        name_col = opts["name_col"] if "name_col" in opts else opts["name-col"]  # compat

        limit_rows = int(opts["limit"] or 0)
        max_new = int(opts["max_new"] or 0)
        dry = bool(opts["dry_run"])
        update_existing = bool(opts["update_existing"])
        log_path = Path(opts["log_file"]) if opts["log_file"] else None

        def log(line: str):
            self.stdout.write(line)
            if log_path:
                log_path.parent.mkdir(parents=True, exist_ok=True)
                with log_path.open("a", encoding="utf-8") as f:
                    f.write(line + "\n")

        wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
        if sheet_name not in wb.sheetnames:
            raise CommandError(f"La hoja '{sheet_name}' no existe. Disponibles: {wb.sheetnames}")
        ws = wb[sheet_name]

        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header_row:
            raise CommandError("La hoja está vacía (no hay encabezado).")

        header = [str(x).strip() if x is not None else "" for x in header_row]
        if name_col not in header:
            raise CommandError(f"No encontré la columna '{name_col}'. Encabezados: {header}")
        idx = header.index(name_col)

        log(f"==> XLSX     : {xlsx}")
        log(f"==> Sheet    : {sheet_name}")
        log(f"==> Column   : {name_col}")
        log(f"==> limit    : {limit_rows} (filas leídas; 0=sin límite)")
        log(f"==> max-new  : {max_new} (nuevos reales; 0=sin límite)")
        log(f"==> update   : {update_existing}")
        log(f"==> dry-run  : {dry}")
        if log_path:
            log(f"==> log-file : {log_path}")

        # 1) leer nombres únicos (en orden de aparición)
        seen = set()
        names = []
        read_rows = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            if limit_rows and read_rows >= limit_rows:
                break
            read_rows += 1

            cell = row[idx] if idx < len(row) else None
            if cell is None:
                continue
            name = str(cell).strip()
            if not name or name in seen:
                continue
            seen.add(name)
            names.append(name)

        log(f"==> Nombres únicos leídos: {len(names)}")

        created = updated = skipped = nomatch = 0
        inserted_new = 0

        # 2) upsert transaccional
        with transaction.atomic():
            for q in names:
                if max_new and inserted_new >= max_new:
                    log(f"[STOP] alcanzado max-new={max_new}")
                    break

                rec = resolve_name_to_taxon(q)
                if not rec:
                    nomatch += 1
                    log(f"[NO_MATCH] {q}")
                    continue

                taxid = int(rec.taxid)
                sci = rec.scientific_name or q
                rank = rec.rank or ""

                obj = Taxon.objects.filter(taxid=taxid).only("taxid").first()
                if obj and not update_existing:
                    skipped += 1
                    log(f"[SKIP_EXISTING] taxid={taxid} name='{sci}'")
                    continue

                # Persistencia mínima (CoL llenará taxonomía real)
                defaults = {
                    "scientific_name": sci,
                    
                }

                # Si tu Taxon tiene rank, lo guardamos; si no, lo ignoramos sin romper.
                if hasattr(Taxon, "rank"):
                    defaults["rank"] = rank

                if obj is None:
                    if dry:
                        created += 1
                        inserted_new += 1
                        log(f"[DRY CREATE] taxid={taxid} name='{sci}'" + (f" rank={rank}" if "rank" in defaults else ""))

                    else:
                        Taxon.objects.create(taxid=taxid, **defaults)
                        created += 1
                        inserted_new += 1
                        log(f"[CREATE] taxid={taxid} name='{sci}'" + (f" rank={rank}" if "rank" in defaults else ""))

                else:
                    if dry:
                        updated += 1
                        log(f"[DRY UPDATE] taxid={taxid} name='{sci}'" + (f" rank={rank}" if "rank" in defaults else ""))
                    else:
                        Taxon.objects.filter(taxid=taxid).update(**defaults)
                        updated += 1
                        log(f"[UPDATE] taxid={taxid} name='{sci}'" + (f" rank={rank}" if "rank" in defaults else ""))

        log("==> RESUMEN")
        log(f"  created      : {created}")
        log(f"  updated      : {updated}")
        log(f"  skipped      : {skipped}")
        log(f"  no_match     : {nomatch}")
        log(f"  inserted_new : {inserted_new}")

        if dry:
            log("==> DRY-RUN: no se escribió nada.")
        else:
            log(str(self.style.SUCCESS("Importación completada.")))
