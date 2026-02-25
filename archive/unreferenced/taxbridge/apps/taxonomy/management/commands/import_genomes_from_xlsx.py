# taxonomy/management/commands/import_genomes_from_xlsx.py
"""
Import genomes from the metazoa_genomes_taxonomy XLSX file.
Create or update NCBIGenome records linked to existing Taxon.
"""
from __future__ import annotations

from pathlib import Path
import openpyxl

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from apps.taxonomy.models import NCBIGenome, Taxon


class Command(BaseCommand):
    help = "Import genomes from XLSX (metazoa_genomes_taxonomy.xlsx)"

    def add_arguments(self, parser):
        parser.add_argument(
            "--file",
            default="data/metazoa_genomes_taxonomy.xlsx",
            help="Path to the XLSX file",
        )
        parser.add_argument("--sheet", default="Sheet1", help="Sheet name")
        parser.add_argument("--dry-run", action="store_true", help="Do not write to DB")
        parser.add_argument("--update", action="store_true", help="Update existing records")

    def handle(self, *args, **opts):
        xlsx = Path(opts["file"])
        if not xlsx.exists():
            raise CommandError(f"File does not exist: {xlsx}")

        dry = opts["dry_run"]
        update = opts["update"]
        sheet_name = opts["sheet"]

        wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
        if sheet_name not in wb.sheetnames:
            raise CommandError(f"Sheet '{sheet_name}' does not exist. Available: {wb.sheetnames}")
        ws = wb[sheet_name]

        # Read headers
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header_row:
            raise CommandError("Sheet is empty or missing header row")

        header = [str(x).strip() if x else "" for x in header_row]
        self.stdout.write(f"Columns found: {header}")

        # Column mapping: XLSX column name -> NCBIGenome field
        col_map = {
            "Phylum": "phylum",
            "Clase": "class_name",
            "Organism": "organism_name",
            "Accession": "accession",
            "Nombre científico": "scientific_name",
            "Nombre común": "common_name",
            "Nombre directorio": "directory_name",
            "Genome level": "genome_level",
            "Genome coverage": "genome_coverage",
            "Contig N50 (kb)": "contig_n50_kb",
            "Scaffold N50 (kb)": "scaffold_n50_kb",
            "Genes": "genes",
            "Protein-Coding": "protein_coding",
        }

        # Find indices
        idx = {}
        for xlsx_col, field in col_map.items():
            if xlsx_col in header:
                idx[field] = header.index(xlsx_col)

        if "accession" not in idx:
            raise CommandError("Columna 'Accession' requerida no encontrada")

        stats = {"created": 0, "updated": 0, "skipped": 0, "errors": 0, "matched_taxon": 0}

        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.stdout.write(f"Procesando {len(rows)} filas...")

        with transaction.atomic():
            for i, row in enumerate(rows, start=2):
                try:
                    accession = self._get_cell(row, idx.get("accession"))
                    if not accession:
                        continue

                    organism = self._get_cell(row, idx.get("organism_name")) or ""

                    # Buscar si existe
                    genome = NCBIGenome.objects.filter(accession=accession).first()
                    
                    if genome and not update:
                        stats["skipped"] += 1
                        continue

                    # Search Taxon by name (case-insensitive)
                    taxon = None
                    if organism:
                        taxon = Taxon.objects.filter(scientific_name__iexact=organism).first()
                        if taxon:
                            stats["matched_taxon"] += 1

                    data = {
                        "organism_name": organism,
                        "phylum": self._get_cell(row, idx.get("phylum")) or "",
                        "class_name": self._get_cell(row, idx.get("class_name")) or "",
                        "common_name": self._get_cell(row, idx.get("common_name")) or "",
                        "directory_name": self._get_cell(row, idx.get("directory_name")) or "",
                        "genome_level": self._get_cell(row, idx.get("genome_level")) or "",
                        "genome_coverage": self._get_float(row, idx.get("genome_coverage")),
                        "contig_n50_kb": self._get_float(row, idx.get("contig_n50_kb")),
                        "scaffold_n50_kb": self._get_float(row, idx.get("scaffold_n50_kb")),
                        "genes": self._get_int(row, idx.get("genes")),
                        "protein_coding": self._get_int(row, idx.get("protein_coding")),
                        "taxon": taxon,
                        "col_match_status": "matched" if taxon else "unmatched",
                    }
                    # Only keep valid NCBIGenome fields
                    valid_fields = {f.name for f in NCBIGenome._meta.get_fields()}
                    data = {k: v for k, v in data.items() if k in valid_fields}

                    if dry:
                        self.stdout.write(f"  [DRY] {accession}: {organism}")
                        stats["created" if not genome else "updated"] += 1
                        continue

                    if genome:
                        for k, v in data.items():
                            setattr(genome, k, v)
                        genome.save()
                        stats["updated"] += 1
                    else:
                        NCBIGenome.objects.create(accession=accession, **data)
                        stats["created"] += 1

                except Exception as e:
                    stats["errors"] += 1
                    self.stderr.write(f"  Error fila {i}: {e}")

        self.stdout.write(self.style.SUCCESS(
            f"\nResultados: {stats['created']} creados, {stats['updated']} actualizados, "
            f"{stats['skipped']} omitidos, {stats['errors']} errores, "
            f"{stats['matched_taxon']} vinculados a Taxon"
        ))

    def _get_cell(self, row, idx):
        if idx is None or idx >= len(row):
            return None
        val = row[idx]
        return str(val).strip() if val is not None else None

    def _get_float(self, row, idx):
        val = self._get_cell(row, idx)
        if not val:
            return None
        try:
            return float(val)
        except (ValueError, TypeError):
            return None

    def _get_int(self, row, idx):
        val = self._get_cell(row, idx)
        if not val:
            return None
        try:
            return int(float(val))
        except (ValueError, TypeError):
            return None
